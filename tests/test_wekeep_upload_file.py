from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

import wekeep_upload_file as export


class WeKeepUploadFileTests(unittest.TestCase):
    @staticmethod
    def _template(path: Path, columns: int, rows: int = 5) -> None:
        workbook = Workbook()
        sheet = workbook.active
        for column in range(1, columns + 1):
            sheet.cell(1, column).value = f"H{column}"
        sheet.cell(rows, columns).value = None
        sheet.row_dimensions[rows].height = 20
        workbook.save(path)

    @staticmethod
    def _row() -> dict:
        return {
            "state": "ready", "order_number": "O-1", "sku_no": "SKU-1",
            "wekeep_product_name": "제품", "source_product_name": "원본", "options": "블루",
            "quantity": 2, "recipient": "홍길동", "phone": "010-1234-5678",
            "zipcode": "01234", "address": "서울시", "message": "문 앞",
        }

    def test_creates_b2c_official_layout(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            template = Path(folder) / "template.xlsx"
            target = Path(folder) / "result.xlsx"
            self._template(template, 15)
            with patch.object(export, "bundled_template_path", return_value=template):
                export.create_wekeep_upload([self._row()], "b2c", target)
            row = [cell.value for cell in load_workbook(target).active[2]]
        self.assertEqual(row[:5], ["O-1", "SKU-1", "제품", "블루", "2"])
        self.assertEqual(row[7:13], ["홍길동", "010-1234-5678", "010-1234-5678", "01234", "서울시", "문 앞"])

    def test_creates_b2b_official_layout(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            template = Path(folder) / "template.xlsx"
            target = Path(folder) / "result.xlsx"
            self._template(template, 29)
            with patch.object(export, "bundled_template_path", return_value=template):
                export.create_wekeep_upload([self._row()], "b2b_buying", target)
            row = [cell.value for cell in load_workbook(target).active[2]]
        self.assertEqual(row[:3], ["O-1", "제품", "SKU-1"])
        self.assertEqual(row[6], "2")
        self.assertEqual(row[9:15], ["홍길동", "010-1234-5678", "010-1234-5678", "01234", "서울시", "문 앞"])

    def test_rejects_review_rows(self) -> None:
        with self.assertRaisesRegex(ValueError, "검토 필요"):
            export.create_wekeep_upload([{"state": "review"}], "b2c")

    def test_rejects_zero_quantity_before_opening_wekeep(self) -> None:
        row = self._row()
        row["quantity"] = 0
        with self.assertRaisesRegex(ValueError, "1 이상의 정수"):
            export.create_wekeep_upload([row], "b2c")


if __name__ == "__main__":
    unittest.main()
