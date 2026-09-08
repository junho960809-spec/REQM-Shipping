from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from excel_loader import load_orders
from shipping_export import export_wekep


class WiselyShippingExportTests(unittest.TestCase):
    def test_wisely_orders_write_channel_to_sales_channel_column(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "wisely.xlsx"
            output = Path(folder) / "shipping.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append([
                "id", "주문번호", "품목주문번호", "상품코드(SKU 코드)", "상품명",
                "상품옵션코드", "상품옵션명", "단가", "수량", "환불수량", "수령인",
                "수령인 전화번호", "우편번호", "주소", "상세주소", "배송메시지",
            ])
            sheet.append([
                "1", "ORDER-1", "LINE-1", "SKU-1", "테스트 상품", "OPT-1", "네이비",
                1000, 1, 0, "홍길동", "010-1234-5678", "01234", "서울시", "101호", "문 앞",
            ])
            workbook.save(source)
            workbook.close()

            orders, _ = load_orders(str(source))
            export_wekep(orders, str(output))
            result = load_workbook(output, read_only=True, data_only=True)
            values = list(result.active.iter_rows(min_row=1, max_row=2, values_only=True))
            result.close()

        self.assertEqual(values[0][1], "판매처")
        self.assertEqual(values[1][1], "와이즐리")


if __name__ == "__main__":
    unittest.main()
