from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from matcher import ProductMatcher
from wekeep_sku_store import load_wekeep_sku_mappings, prepare_wekeep_orders
from wekeep_upload_file import create_wekeep_upload


class WeKeepSetTests(unittest.TestCase):
    def setUp(self) -> None:
        self.order = {
            "status": "exact", "components": "A + B + C", "quantity": 2,
            "order_number": "SET-TEST", "channel": "테스트", "product_name": "본품 + 옵션 세트",
            "recipient": "테스트", "phone": "010-0000-0000", "zipcode": "01234",
            "address": "테스트 주소", "message": "배송 메모",
        }
        self.mappings = [
            {"item_code": code, "sku_no": "SKU-" + code, "product_name": "품목 " + code}
            for code in "ABC"
        ]

    def prepare(self, components: str, **changes) -> list[dict]:
        return prepare_wekeep_orders(
            [{**self.order, "components": components, **changes}],
            order_kind="b2c", mappings=self.mappings,
        )

    def test_plain_and_mixed_sets_keep_every_component(self) -> None:
        for components, expected in [
            ("A", [("A", 2)]),
            (" A + B + C ", [("A", 2), ("B", 2), ("C", 2)]),
            ("A×1 + B + C×2", [("A", 2), ("B", 2), ("C", 4)]),
            ("A + B×2 + C", [("A", 2), ("B", 4), ("C", 2)]),
            ("A×1.0 + B", [("A", 2), ("B", 2)]),
            ("A + B×2 + a", [("A", 4), ("B", 4)]),
        ]:
            with self.subTest(components=components):
                rows = self.prepare(components)
                self.assertEqual([(r["item_code"], r["quantity"]) for r in rows], expected)
                self.assertTrue(all(r["state"] == "ready" for r in rows))
                for row in rows:
                    for key in ("order_number", "recipient", "phone", "zipcode", "address", "message"):
                        self.assertEqual(row[key], self.order[key])

    def test_missing_option_blocks_entire_set(self) -> None:
        rows = self.prepare("A×1 + UNKNOWN + B×2")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["state"], "review")
        self.assertIn("UNKNOWN", rows[0]["reason"])
        self.assertNotIn("quantity", rows[0])

    def test_invalid_component_never_produces_partial_shipment(self) -> None:
        for component in ("B×0", "B×-1", "B×1.5", "B×", "B×bad", "B×1×2", "B×NaN", ""):
            with self.subTest(component=component):
                rows = self.prepare("A×1 + " + component)
                self.assertEqual(len(rows), 1)
                self.assertEqual(rows[0]["state"], "review")
                self.assertNotIn("sku_no", rows[0])

    def test_repeated_items_are_not_merged_across_source_orders(self) -> None:
        rows = prepare_wekeep_orders(
            [{**self.order, "components": "A + A"},
             {**self.order, "components": "A", "order_number": "OTHER"}],
            order_kind="b2c", mappings=self.mappings,
        )
        self.assertEqual([(r["order_number"], r["quantity"]) for r in rows], [("SET-TEST", 4), ("OTHER", 2)])

    def test_reported_battery_case_codes_with_bundled_skus(self) -> None:
        rows = prepare_wekeep_orders(
            [{**self.order, "components": "[RQM]-QP1000C1-WH + CASE-QP1000C-Hand-Re"}],
            order_kind="b2c", mappings=load_wekeep_sku_mappings(local_path=Path("__missing__")),
        )
        self.assertEqual([(r["sku_no"], r["quantity"]) for r in rows],
                         [("66896227366381", 2), ("54476196066367", 2)])
        self.assertTrue(all(r["state"] == "ready" for r in rows))

    def test_name_matching_produces_explicit_component_quantities(self) -> None:
        matcher = ProductMatcher([
            {"item_code": "A", "standard_name": "본품"},
            {"item_code": "B", "standard_name": "옵션"},
        ], [], [])
        matched = matcher.match({"matched_name": "본품 / 옵션"})
        self.assertEqual(matched["status"], "exact")
        self.assertEqual(matched["components"], "A×1 + B×1")
        self.assertEqual(len(self.prepare(matched["components"])), 2)

    def test_sets_reach_all_four_upload_formats_with_multiplied_quantities(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            for kind, quantity_column in [("b2c", 4), ("b2c_buying", 5), ("b2b", 7), ("b2b_buying", 7)]:
                with self.subTest(kind=kind):
                    rows = prepare_wekeep_orders(
                        [{**self.order, "components": "A + B×2 + C"}],
                        order_kind=kind, mappings=self.mappings,
                    )
                    target = create_wekeep_upload(rows, kind, Path(folder) / f"{kind}.xlsx")
                    workbook = load_workbook(target, read_only=True, data_only=True)
                    try:
                        sheet = workbook.active
                        self.assertEqual([sheet.cell(r, quantity_column).value for r in (2, 3, 4)], ["2", "4", "2"])
                        self.assertEqual([sheet.cell(r, 1).value for r in (2, 3, 4)], ["SET-TEST"] * 3)
                    finally:
                        workbook.close()


if __name__ == "__main__":
    unittest.main()
