from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from wekeep_tracking_service import (apply_manual_tracking_number, export_tracking_workbook,
                                     manual_tracking_candidate_indexes, reconcile_tracking_numbers)


class WeKeepTrackingServiceTests(unittest.TestCase):
    @staticmethod
    def _orders() -> list[dict]:
        return [
            {"order_number": "O-1", "channel": "셀메이트", "product_name": "본품", "quantity": 1,
             "recipient": "홍길동", "phone": "010-1234-5678", "zipcode": "01234",
             "address": "서울시 테스트로 1", "message": "문 앞", "serial_number": "S-1"},
            {"order_number": "O-1", "channel": "셀메이트", "product_name": "옵션", "quantity": 1,
             "recipient": "홍길동", "phone": "010-1234-5678", "zipcode": "01234",
             "address": "서울시 테스트로 1", "message": "문 앞", "serial_number": "S-2"},
        ]

    def test_exact_order_and_recipient_fill_same_invoice_on_set_rows(self) -> None:
        results = reconcile_tracking_numbers(self._orders(), [{
            "order_number": "O-1", "recipient": "홍길동", "tracking_number": "5409-8333-2747",
        }])
        self.assertEqual([row["tracking_number"] for row in results], ["540983332747"] * 2)
        self.assertTrue(all(row["tracking_match_state"] == "matched" for row in results))

    def test_recipient_mismatch_never_fills_invoice(self) -> None:
        result = reconcile_tracking_numbers(self._orders()[:1], [{
            "order_number": "O-1", "recipient": "다른사람", "tracking_number": "540983332747",
        }])[0]
        self.assertEqual(result["tracking_match_state"], "review")
        self.assertEqual(result["tracking_number"], "")

    def test_phone_and_address_are_checked_when_wekeep_exposes_them(self) -> None:
        remote = {"order_number": "O-1", "recipient": "홍길동", "phone": "01012345678",
                  "address": "부산시 다른로 2", "tracking_number": "540983332747"}
        result = reconcile_tracking_numbers(self._orders()[:1], [remote])[0]
        self.assertEqual(result["tracking_match_state"], "review")

    def test_combined_shipping_copies_one_invoice_to_distinct_identical_deliveries(self) -> None:
        orders = self._orders()[:1] + [{**self._orders()[0], "order_number": "O-2", "product_name": "다른 주문"}]
        rows = reconcile_tracking_numbers(orders, [
            {"order_number": "O-1", "recipient": "홍길동", "tracking_number": "540983332747"},
            {"order_number": "O-2", "recipient": "홍길동", "tracking_number": "-"},
        ])
        self.assertEqual([row["tracking_number"] for row in rows], ["540983332747"] * 2)
        self.assertEqual(rows[1]["tracking_match_state"], "matched")
        self.assertIn("합포장", rows[1]["tracking_match_reason"])

    def test_combined_shipping_requires_all_four_delivery_fields_to_match(self) -> None:
        base = self._orders()[0]
        for field, changed in (("recipient", "김길동"), ("phone", "010-9999-9999"),
                               ("zipcode", "99999"), ("address", "서울시 다른로 2")):
            with self.subTest(field=field):
                second = {**base, "order_number": "O-2", field: changed}
                rows = reconcile_tracking_numbers([base, second], [
                    {"order_number": "O-1", "recipient": "홍길동", "tracking_number": "540983332747"},
                    {"order_number": "O-2", "recipient": second["recipient"], "tracking_number": "-"},
                ])
                self.assertEqual(rows[1]["tracking_match_state"], "pending")
                self.assertEqual(rows[1]["tracking_number"], "")

    def test_combined_shipping_does_not_choose_between_multiple_invoices(self) -> None:
        base = self._orders()[0]
        orders = [base, {**base, "order_number": "O-2"}, {**base, "order_number": "O-3"}]
        rows = reconcile_tracking_numbers(orders, [
            {"order_number": "O-1", "recipient": "홍길동", "tracking_number": "11111111"},
            {"order_number": "O-2", "recipient": "홍길동", "tracking_number": "22222222"},
            {"order_number": "O-3", "recipient": "홍길동", "tracking_number": "-"},
        ])
        self.assertEqual(rows[2]["tracking_match_state"], "pending")

    def test_manual_invoice_applies_to_the_exact_delivery_group(self) -> None:
        base = self._orders()[0]
        rows = [
            {**base, "tracking_match_state": "pending", "tracking_number": ""},
            {**base, "order_number": "O-2", "tracking_match_state": "pending", "tracking_number": ""},
            {**base, "order_number": "O-3", "address": "서울시 다른로 2", "tracking_match_state": "pending", "tracking_number": ""},
        ]
        updated = apply_manual_tracking_number(rows, 0, "5409-8333-2747")
        self.assertEqual([row["tracking_number"] for row in updated], ["540983332747", "540983332747", ""])
        self.assertTrue(all("수동 입력" in updated[index]["tracking_match_reason"] for index in (0, 1)))

    def test_manual_invoice_does_not_require_delivery_fields(self) -> None:
        rows = [
            {"order_number": "O-1", "product_name": "본품", "tracking_match_state": "pending"},
            {"order_number": "O-1", "product_name": "옵션", "tracking_match_state": "pending"},
            {"order_number": "O-2", "product_name": "별도 주문", "tracking_match_state": "pending"},
        ]
        updated = apply_manual_tracking_number(rows, 0, "540983332747")
        self.assertEqual([row.get("tracking_number", "") for row in updated], ["540983332747", "540983332747", ""])

    def test_manual_targets_can_be_limited_after_worker_review(self) -> None:
        base = self._orders()[0]
        rows = [base, {**base, "order_number": "O-2"}]
        self.assertEqual(manual_tracking_candidate_indexes(rows, 0), [0, 1])
        updated = apply_manual_tracking_number(rows, 0, "540983332747", [0])
        self.assertEqual(updated[0]["tracking_number"], "540983332747")
        self.assertEqual(updated[1].get("tracking_number", ""), "")

    def test_pending_missing_and_multiple_invoices_stay_blank(self) -> None:
        for remote, expected in [
            ([{"order_number": "O-1", "recipient": "홍길동", "tracking_number": "-"}], "pending"),
            ([], "not_found"),
            ([{"order_number": "O-1", "recipient": "홍길동", "tracking_number": "11111111", "additional_tracking": "22222222"}], "review"),
        ]:
            with self.subTest(expected=expected):
                result = reconcile_tracking_numbers(self._orders()[:1], remote)[0]
                self.assertEqual(result["tracking_match_state"], expected)
                self.assertEqual(result["tracking_number"], "")

    def test_export_preserves_11_column_layout_and_writes_invoice_as_text(self) -> None:
        rows = reconcile_tracking_numbers(self._orders(), [{
            "order_number": "O-1", "recipient": "홍길동", "tracking_number": "540983332747",
        }])
        with tempfile.TemporaryDirectory() as folder:
            path = export_tracking_workbook(rows, Path(folder) / "tracking.xlsx")
            workbook = load_workbook(path, data_only=True)
            try:
                sheet = workbook.active
                self.assertEqual(sheet.max_column, 11)
                self.assertEqual([cell.value for cell in sheet[1]], [
                    "주문번호", "판매처", "상품명", "수량", "수령자", "핸드폰", "우편번호",
                    "주소", "배송메세지", "송장번호", "일련번호",
                ])
                self.assertEqual([sheet.cell(row, 10).value for row in (2, 3)], ["540983332747"] * 2)
                self.assertEqual([sheet.cell(row, 11).value for row in (2, 3)], ["S-1", "S-2"])
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
