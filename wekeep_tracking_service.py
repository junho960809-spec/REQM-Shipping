"""Fetch WeKeep tracking numbers and reconcile them with loaded REQM orders."""
from __future__ import annotations

import re
from collections import defaultdict
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from integration_credential_store import load_integration_credentials
from shipping_export import export_wekep
from wekeep_order_automation import ensure_authenticated
from wekeep_report_service import PROFILE_PATH, launch_wekeep_context


ORDER_SEARCH_URL = "https://fbw.wekeep.co.kr/fbw/admin/v2/order/searchOrder"
SALE_CHANNELS = {
    "b2c": ("SCNC75B2282BEYUXP", "주식회사 리큐엠(B2C)"),
    "b2c_buying": ("SCN2F09B6F8C7YUKO", "주식회사 리큐엠(사입형B2C)"),
    "b2b": ("SCN477724148DY0CS", "주식회사 리큐엠(B2B)"),
}


def _compact(value: object) -> str:
    return re.sub(r"[^0-9a-z가-힣]", "", str(value or "").casefold())


def _phone(value: object) -> str:
    return re.sub(r"\D", "", str(value or ""))


def _tracking_numbers(value: object) -> tuple[str, ...]:
    return tuple(dict.fromkeys(re.findall(r"[0-9]{8,20}", str(value or "").replace("-", ""))))


def collect_tracking_rows(page, registered_date: date, sale_channel: tuple[str, str]) -> list[dict]:
    """Search one WeKeep sales channel/date and return every visible result page."""
    page.goto(ORDER_SEARCH_URL, wait_until="domcontentloaded", timeout=60_000)
    credentials = load_integration_credentials()
    ensure_authenticated(page, credentials.get("wekeep_user_id", ""), credentials.get("wekeep_password", ""))
    if "/order/searchOrder" not in page.url:
        page.goto(ORDER_SEARCH_URL, wait_until="domcontentloaded", timeout=60_000)
    page.wait_for_function("() => document.querySelectorAll('#saleChannel_search option').length > 1", timeout=15_000)
    day = registered_date.isoformat()
    # This page reloads from select/date change handlers. Set every filter and submit
    # once in page JavaScript so a reload cannot discard half of the criteria.
    try:
        page.evaluate(r"""({day, saleChannelValue, saleChannelLabel}) => {
          const sale = document.querySelector('#saleChannel_search');
          const option = [...sale.options].find(item => item.value === saleChannelValue);
          if (!option) throw new Error('요청한 위킵 판매처를 찾지 못했습니다: ' + saleChannelLabel);
          sale.value = option.value;
          document.querySelector('#startDate_search').value = day;
          document.querySelector('#endDate_search').value = day;
          const length = document.querySelector('#searchLength');
          if (length && [...length.options].some(item => item.textContent.includes('100개'))) {
            length.value = [...length.options].find(item => item.textContent.includes('100개')).value;
          }
          if (typeof searchDetailOrderList !== 'function') throw new Error('위킵 상세검색 함수를 찾지 못했습니다.');
          searchDetailOrderList();
        }""", {"day": day, "saleChannelValue": sale_channel[0], "saleChannelLabel": sale_channel[1]})
    except Exception as exc:
        if "navigation" not in str(exc).casefold() and "context was destroyed" not in str(exc).casefold():
            raise
    page.wait_for_timeout(1_200)
    rows: list[dict] = []
    seen_pages: set[str] = set()
    while True:
        page_rows = page.evaluate(r"""() => {
          const clean = value => String(value || '').replace(/\s+/g, ' ').trim();
          const table = [...document.querySelectorAll('table')].find(t => {
            const headers = [...t.querySelectorAll('th')].map(th => clean(th.textContent));
            return headers.includes('판매처주문번호') && headers.includes('송장번호');
          });
          if (!table) throw new Error('위킵 주문 상세검색 결과 표를 찾지 못했습니다.');
          const headers = [...table.querySelectorAll('th')].map(th => clean(th.textContent));
          const index = name => headers.indexOf(name);
          return [...table.querySelectorAll('tbody tr')].map(tr => {
            const cells = [...tr.querySelectorAll('td')].map(td => clean(td.textContent));
            return {registered_date: cells[index('주문등록일')] || '', order_number: cells[index('판매처주문번호')] || '', recipient: cells[index('수령자')] || '', tracking_number: cells[index('송장번호')] || '', additional_tracking: cells[index('추가송장')] || '', order_status: cells[index('주문상태')] || ''};
          }).filter(row => row.order_number || row.recipient);
        }""")
        signature = "|".join(str(row) for row in page_rows)
        if signature in seen_pages:
            break
        seen_pages.add(signature)
        rows.extend(page_rows)
        next_link = page.get_by_role("link", name=">", exact=True)
        if next_link.count() != 1 or not next_link.is_visible():
            break
        next_link.evaluate("link => link.click()")
        page.wait_for_timeout(800)
    return rows


def fetch_wekeep_tracking_rows(registered_date: date, order_kind: str = "b2c") -> list[dict]:
    from playwright.sync_api import sync_playwright
    sale_channel = SALE_CHANNELS.get(order_kind)
    if not sale_channel:
        raise ValueError(f"송장 조회를 지원하지 않는 판매처 유형입니다: {order_kind}")
    PROFILE_PATH.mkdir(parents=True, exist_ok=True)
    with sync_playwright() as playwright:
        context = launch_wekeep_context(playwright, headless=True)
        try:
            page = context.pages[0] if context.pages else context.new_page()
            return collect_tracking_rows(page, registered_date, sale_channel)
        finally:
            context.close()


def reconcile_tracking_numbers(orders: list[dict], remote_rows: list[dict]) -> list[dict]:
    """Match conservatively; one uncertain remote order never fills an invoice."""
    by_order: dict[str, list[dict]] = defaultdict(list)
    for row in remote_rows:
        by_order[_compact(row.get("order_number"))].append(row)
    results: list[dict] = []
    for order in orders:
        order_number = _compact(order.get("order_number"))
        candidates = list(by_order.get(order_number, [])) if order_number else []
        if not candidates:
            results.append({**order, "tracking_match_state": "not_found", "tracking_match_reason": "위킵에서 주문번호를 찾지 못했습니다.", "tracking_number": ""})
            continue
        recipient = _compact(order.get("recipient"))
        recipient_matches = [row for row in candidates if not recipient or _compact(row.get("recipient")) == recipient]
        if recipient_matches:
            candidates = recipient_matches
        elif recipient:
            results.append({**order, "tracking_match_state": "review", "tracking_match_reason": "주문번호는 같지만 수령인이 다릅니다.", "tracking_number": ""})
            continue
        for key, normalizer in (("phone", _phone), ("address", _compact)):
            expected = normalizer(order.get(key))
            available = [row for row in candidates if normalizer(row.get(key))]
            if expected and available:
                candidates = [row for row in available if normalizer(row.get(key)) == expected]
                if not candidates:
                    break
        if not candidates:
            results.append({**order, "tracking_match_state": "review", "tracking_match_reason": "전화번호 또는 주소가 다릅니다.", "tracking_number": ""})
            continue
        tracking = tuple(dict.fromkeys(number for row in candidates for number in (*_tracking_numbers(row.get("tracking_number")), *_tracking_numbers(row.get("additional_tracking")))))
        if not tracking:
            results.append({**order, "tracking_match_state": "pending", "tracking_match_reason": "주문은 확인됐지만 송장이 아직 발급되지 않았습니다.", "tracking_number": ""})
        elif len(tracking) > 1:
            results.append({**order, "tracking_match_state": "review", "tracking_match_reason": "한 주문에 여러 송장이 확인됐습니다: " + ", ".join(tracking), "tracking_number": ""})
        else:
            results.append({**order, "tracking_match_state": "matched", "tracking_match_reason": "주문번호·수령인 대조 완료", "tracking_number": tracking[0]})
    return results


def export_tracking_workbook(rows: list[dict], output_path: str | Path) -> Path:
    """Write the existing 11-column carrier workbook with invoice numbers in column J."""
    target = Path(output_path)
    export_wekep(rows, str(target))
    workbook = load_workbook(target)
    try:
        sheet = workbook.active
        for row_index, row in enumerate(rows, start=2):
            sheet.cell(row_index, 10).value = str(row.get("tracking_number") or "")
            sheet.cell(row_index, 10).number_format = "@"
        workbook.save(target)
    finally:
        workbook.close()
    return target
